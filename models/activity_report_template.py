# -*- coding: utf-8 -*-
"""Excel templates for activity reports.

A template is a workbook plus a list of mappings, as in expense_scan: a cell
("B3") receives a header value once, a column letter ("A") receives a value
on every line of a block. On top of that, an activity report is a month
grid: one column per day, from ``first_day_column`` on, and two blocks of
lines, missions then time off.

The template's blocks are resized to exactly the report's lines; formulas
below a block (totals) follow. The day columns of days the month does not
have (29 to 31) are emptied and hidden.

Users download the blank workbook, lay it out in Excel or LibreOffice (logo,
colours, extra rows or columns), then upload it back.
"""
import base64
import io
import logging
import re
import zipfile
from copy import copy

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

#: Values a line column can receive.
LINE_VALUES = [
    ('label', "Client, or time off type"),
    ('client', "Client"),
    ('project', "Mission"),
    ('leave_type', "Time off type"),
    ('total', "Line total (days)"),
]
#: Values a header cell can receive.
HEADER_VALUES = [
    ('employee', "Employee"),
    ('company', "Company"),
    ('month', "Month (first day, as a date)"),
    ('month_label', "Month (text)"),
    ('potential_days', "Working days"),
    ('status', "Status"),
    ('submit_date', "Submission date"),
    ('validate_date', "Validation date"),
    ('client', "Client (client report)"),
    ('mission_total', "Mission days"),
    ('absence_total', "Time off days"),
    ('total', "Total days"),
]
VALUES = LINE_VALUES + [item for item in HEADER_VALUES
                        if item[0] not in {key for key, _label in LINE_VALUES}]
CELL_RE = re.compile(r"^([A-Z]{1,3})([0-9]*)$")
#: Reference ("$L$70") or range ("L8:L68") in a formula of the same sheet.
CELL_REF_RE = re.compile(
    r"(?<![A-Za-z0-9_!'.$])(\$?[A-Z]{1,3}\$?)(\d+)"
    r"(?::(\$?[A-Z]{1,3}\$?)(\d+))?(?![0-9A-Za-z_(])")
MAX_DAYS = 31


def _import_openpyxl():
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError as error:
        raise UserError(_("The openpyxl library is missing on the server.")) from error
    return openpyxl


class MissionActivityReportTemplate(models.Model):
    _name = 'mission.activity.report.template'
    _description = "Activity report Excel template"
    _order = 'sequence, name'

    name = fields.Char(required=True, translate=True)
    sequence = fields.Integer(default=10)
    active = fields.Boolean(default=True)
    company_id = fields.Many2one('res.company', string="Company")
    kind = fields.Selection(
        [('internal', "Internal (missions and time off)"), ('client', "Client (missions of one client)")],
        required=True, default='internal',
        help="A client report is exported once per client of the month.")
    file = fields.Binary(string="Excel file", attachment=True)
    filename = fields.Char()
    file_generated = fields.Boolean(
        readonly=True, copy=False,
        help="The workbook is the module's blank one: module updates remake it.")
    sheet_name = fields.Char(string="Sheet", help="Empty: the first sheet of the workbook.")
    first_day_column = fields.Char(
        string="Day 1 column", required=True, default='D',
        help="Column of the 1st of the month; the next 30 columns hold the other days.")
    week_row = fields.Integer(string="Week numbers row", help="0: none.")
    day_row = fields.Integer(string="Day numbers row", help="0: none.")
    weekday_row = fields.Integer(string="Weekdays row", help="0: none.")
    mission_first_row = fields.Integer(string="First mission row", required=True, default=10)
    mission_last_row = fields.Integer(string="Last mission row", required=True, default=12)
    absence_first_row = fields.Integer(
        string="First time off row", help="0: no time off block (client reports).")
    absence_last_row = fields.Integer(string="Last time off row")
    color_days = fields.Boolean(
        string="Colour the days", default=True,
        help="Fills the day columns of weekends, public holidays and time off.")
    weekend_color = fields.Char(default='D9D9D9')
    holiday_color = fields.Char(default='FFF2CC')
    absence_color = fields.Char(default='C6EFCE')
    cell_ids = fields.One2many(
        'mission.activity.report.template.cell', 'template_id', string="Mappings", copy=True)
    note = fields.Text()

    @api.constrains('mission_first_row', 'mission_last_row', 'absence_first_row',
                    'absence_last_row', 'first_day_column')
    def _check_layout(self):
        for template in self:
            if template.mission_first_row < 1 or template.mission_last_row < template.mission_first_row:
                raise UserError(_("The mission rows of template “%s” are inconsistent.", template.name))
            if template.absence_first_row and (
                    template.absence_last_row < template.absence_first_row
                    or template.absence_first_row <= template.mission_last_row):
                raise UserError(_("The time off rows of template “%s” must come after the mission rows.",
                                  template.name))
            if not CELL_RE.match((template.first_day_column or '').strip().upper()) \
                    or CELL_RE.match(template.first_day_column.strip().upper()).group(2):
                raise UserError(_("“Day 1 column” takes a column letter, such as D."))

    # ------------------------------------------------------------------
    # Blank workbook
    # ------------------------------------------------------------------

    def action_generate_file(self):
        """Blank workbook laid out after the template's settings and mappings."""
        for template in self:
            template.with_context(mission_report_blank_file=True).write({
                'file': base64.b64encode(template._blank_workbook()),
                'filename': "%s.xlsx" % re.sub(r'[\\/:*?"<>|]+', '-', template.name),
                'file_generated': True,
            })
        return True

    def write(self, vals):
        # A workbook uploaded by the user is theirs: module updates must not
        # replace it with a blank one.
        if 'file' in vals and not self.env.context.get('mission_report_blank_file'):
            vals['file_generated'] = False
        return super().write(vals)

    @api.model
    def _fill_blank_files(self):
        """(Re)make the blank workbook of the shipped templates.

        Run on every module update: a template whose workbook was never
        replaced by the user follows the module's layout and the company's
        language.
        """
        try:
            import openpyxl  # noqa: F401, PLC0415
        except ImportError:
            _logger.warning("openpyxl missing: activity report templates shipped without workbook")
            return
        lang = self.env.company.partner_id.lang or self.env.user.lang or 'en_US'
        for xmlid in ('mission_report.activity_report_template_internal',
                      'mission_report.activity_report_template_client'):
            template = self.env.ref(xmlid, raise_if_not_found=False)
            if template and (not template.file or template.file_generated):
                template.with_context(lang=lang).action_generate_file()

    def _blank_labels(self):
        """Labels of the blank workbook, in the context's language."""
        header = {
            'employee': _("Employee"),
            'company': _("Company"),
            'month': _("Month"),
            'month_label': _("Month"),
            'potential_days': _("Working days"),
            'status': _("Status"),
            'submit_date': _("Submission date"),
            'validate_date': _("Validation date"),
            'client': _("Client"),
            'mission_total': _("Mission days"),
            'absence_total': _("Time off days"),
            'total': _("Total days"),
        }
        line = {
            'label': _("Client"),
            'client': _("Client"),
            'project': _("Mission"),
            'leave_type': _("Time off type"),
            'total': _("Total"),
        }
        return header, line

    def _blank_workbook(self):
        """A month grid laid out like the PDF report.

        Title, header values with their labels, a colour legend, the day
        header rows, the mission and time off blocks, a total line and two
        signature boxes; A4 landscape, one page wide.
        """
        self.ensure_one()
        openpyxl = _import_openpyxl()
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: PLC0415
        from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: PLC0415

        header_labels, line_labels = self._blank_labels()
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = (self.sheet_name or _("Activity"))[:31]

        font_name = 'Calibri'
        base = Font(name=font_name, size=9)
        bold = Font(name=font_name, size=9, bold=True)
        thin = Side(style='thin', color='A6A6A6')
        medium = Side(style='medium', color='595959')
        box = Border(left=thin, right=thin, top=thin, bottom=thin)
        head_fill = PatternFill('solid', start_color='E7EEF6')
        section_fill = PatternFill('solid', start_color='F2F2F2')
        center = Alignment(horizontal='center', vertical='center')
        left = Alignment(horizontal='left', vertical='center', indent=1)
        right = Alignment(horizontal='right', vertical='center')

        first_day = column_index_from_string(self.first_day_column.strip().upper())
        last_column = first_day + MAX_DAYS - 1
        day_columns = range(first_day, last_column + 1)
        # One mapping per column, the missions' one first: a column can hold
        # the client for missions and the time off type below.
        line_columns = {}
        for cell_map in self.cell_ids.filtered(lambda c: c.kind == 'line' and c.value):
            column = column_index_from_string(cell_map.cell.strip().upper())
            if column not in line_columns or line_columns[column].block == 'absence':
                line_columns[column] = cell_map
        label_columns = range(1, first_day)

        # Title, across the whole grid.
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
        title = sheet.cell(row=1, column=1, value=_("Monthly activity report"))
        title.font = Font(name=font_name, size=14, bold=True)
        title.alignment = center
        sheet.row_dimensions[1].height = 24

        # Header: each value, and its label in the cell on its left.
        header_rows = set()
        for cell_map in self.cell_ids.filtered(lambda c: c.kind == 'header' and c.value):
            cell = sheet[cell_map.cell.strip().upper()]
            header_rows.add(cell.row)
            cell.font = Font(name=font_name, size=10)
            cell.alignment = Alignment(horizontal='left', vertical='center')
            if cell.column > 1:
                label = sheet.cell(row=cell.row, column=cell.column - 1,
                                   value="%s :" % header_labels[cell_map.value]
                                   if self.env.lang and self.env.lang.startswith('fr')
                                   else "%s:" % header_labels[cell_map.value])
                label.font = Font(name=font_name, size=10, bold=True)
                label.alignment = right
            if cell_map.value in ('month', 'submit_date', 'validate_date'):
                cell.number_format = 'DD/MM/YYYY'

        # Colour legend, just above the grid when that row is free.
        grid_rows = [row for row in (self.week_row, self.day_row, self.weekday_row,
                                     self.mission_first_row - 1) if row]
        legend_row = min(grid_rows) - 1
        if self.color_days and legend_row > 1 and legend_row not in header_rows:
            legend = [(self.weekend_color, _("Weekend")), (self.holiday_color, _("Public holiday"))]
            if self.kind == 'internal':
                legend.append((self.absence_color, _("Time off")))
            column = first_day
            for color, text in legend:
                swatch = sheet.cell(row=legend_row, column=column)
                swatch.fill = PatternFill('solid', start_color=color)
                swatch.border = box
                label = sheet.cell(row=legend_row, column=column + 1, value=text)
                label.font = base
                column += 6

        # Day header rows: week numbers, day numbers, weekdays.
        for row, title_text in ((self.week_row, _("Week")), (self.day_row, _("Day")),
                                (self.weekday_row, '')):
            if not row:
                continue
            if first_day > 1 and title_text:
                label = sheet.cell(row=row, column=first_day - 1, value=title_text)
                label.font = bold
                label.alignment = right
            for column in day_columns:
                cell = sheet.cell(row=row, column=column)
                cell.font = bold
                cell.alignment = center
                cell.border = box
                cell.fill = head_fill

        # Column titles above the mission lines.
        title_row = self.mission_first_row - 1
        if title_row >= 1:
            for column, cell_map in line_columns.items():
                if cell_map.block == 'absence':
                    continue
                head = sheet.cell(row=title_row, column=column, value=line_labels[cell_map.value])
                head.font = bold
                head.fill = head_fill
                head.border = box
                head.alignment = center

        blocks = [('mission', self.mission_first_row, self.mission_last_row)]
        if self.absence_first_row:
            section_row = self.absence_first_row - 1
            for column in range(1, last_column + 1):
                cell = sheet.cell(row=section_row, column=column)
                cell.fill = section_fill
                cell.border = box
            section = sheet.cell(row=section_row, column=1, value=_("Time off"))
            section.font = bold
            section.alignment = left
            blocks.append(('absence', self.absence_first_row, self.absence_last_row))
        for block, first, last in blocks:
            for row in range(first, last + 1):
                for column in list(label_columns) + list(day_columns):
                    cell = sheet.cell(row=row, column=column)
                    cell.border = box
                    cell.font = base
                    cell_map = line_columns.get(column)
                    if column in day_columns or (cell_map and cell_map.value == 'total'):
                        cell.alignment = center
                        cell.number_format = '0.##'
                    else:
                        cell.alignment = left
                    if cell_map and cell_map.value == 'total':
                        cell.font = bold

        # Total line: day by day, sum of the blocks above.
        total_row = blocks[-1][2] + 1
        top = Border(left=thin, right=thin, top=medium, bottom=thin)
        for column in list(label_columns) + list(day_columns):
            cell = sheet.cell(row=total_row, column=column)
            cell.border = top
            cell.fill = head_fill
            cell.font = bold
            cell.alignment = center
            cell.number_format = '0.##'
        sheet.cell(row=total_row, column=1, value=_("Total")).alignment = left
        for column in day_columns:
            letter = get_column_letter(column)
            ranges = ["%s%s:%s%s" % (letter, first, letter, last) for _block, first, last in blocks]
            sheet.cell(row=total_row, column=column).value = \
                "=%s" % "+".join("SUM(%s)" % item for item in ranges)
        for column, cell_map in line_columns.items():
            if cell_map.value == 'total':
                letter = get_column_letter(column)
                ranges = ["%s%s:%s%s" % (letter, first, letter, last) for _block, first, last in blocks]
                sheet.cell(row=total_row, column=column).value = \
                    "=%s" % "+".join("SUM(%s)" % item for item in ranges)

        # Signature boxes, two rows below the total.
        box_top = total_row + 2
        boxes = [(1, max(first_day - 2, 1), _("Manager signature")),
                 (first_day + 8, first_day + 20, _("Employee signature"))]
        for start, end, text in boxes:
            sheet.merge_cells(start_row=box_top, start_column=start, end_row=box_top + 5, end_column=end)
            for row in range(box_top, box_top + 6):
                for column in range(start, end + 1):
                    sheet.cell(row=row, column=column).border = box
            cell = sheet.cell(row=box_top, column=start, value=text)
            cell.font = bold
            cell.alignment = Alignment(horizontal='left', vertical='top', indent=1)

        # Widths, frozen panes, printing.
        for column in label_columns:
            cell_map = line_columns.get(column)
            width = 10 if cell_map and cell_map.value == 'total' else 28
            sheet.column_dimensions[get_column_letter(column)].width = width
        for column in day_columns:
            sheet.column_dimensions[get_column_letter(column)].width = 4.3
        sheet.freeze_panes = sheet.cell(row=self.mission_first_row, column=first_day)
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins.left = sheet.page_margins.right = 0.4
        sheet.page_margins.top = sheet.page_margins.bottom = 0.5
        sheet.print_options.horizontalCentered = True
        sheet.print_area = "A1:%s%s" % (get_column_letter(last_column), box_top + 5)

        output = io.BytesIO()
        book.save(output)
        return output.getvalue()

    # ------------------------------------------------------------------
    # Export
    # ------------------------------------------------------------------

    def _export(self, report, partner_id=None):
        """The template's workbook, filled with one report."""
        self.ensure_one()
        openpyxl = _import_openpyxl()
        from openpyxl.styles import PatternFill  # noqa: PLC0415
        from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: PLC0415
        if not self.file:
            raise UserError(_("Template “%s” has no Excel file.", self.name))

        data = report._get_render_data(partner_id=partner_id if self.kind == 'client' else None)
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.file)))
        sheet = book[self.sheet_name] if self.sheet_name else book.worksheets[0]
        first_day = column_index_from_string(self.first_day_column.strip().upper())
        days = data['days']

        blocks = [('mission', self.mission_first_row, self.mission_last_row, data['missions'])]
        if self.absence_first_row:
            blocks.append(('absence', self.absence_first_row, self.absence_last_row,
                           data['absences'] if self.kind == 'internal' else []))
        # Resize bottom-up: resizing a block leaves the rows above it in
        # place. A block empty of lines keeps one blank row.
        for _block, first, last, lines in sorted(blocks, key=lambda item: -item[1]):
            self._resize_block(sheet, first, last, max(len(lines), 1))
        # Then fill top-down, each block pushed by the growth of those above.
        filled_rows = []
        shift = 0
        for block, first, last, lines in blocks:
            count = max(len(lines), 1)
            size = last - first + 1
            first += shift
            shift += count - size
            columns = {column_index_from_string(c.cell.strip().upper()): c.value
                       for c in self.cell_ids
                       if c.kind == 'line' and c.value and c.block in (block, 'both')}
            for index in range(count):
                row = first + index
                line = lines[index] if index < len(lines) else None
                for column, value in columns.items():
                    sheet.cell(row=row, column=column).value = \
                        self._line_value(line, value) if line else None
                for offset in range(MAX_DAYS):
                    cell = sheet.cell(row=row, column=first_day + offset)
                    cell.value = line['values'][offset] if line and offset < len(days) else None
            filled_rows.append((first, first + count - 1))

        for cell_map in self.cell_ids.filtered(lambda c: c.kind == 'header' and c.value):
            sheet[cell_map.cell.strip().upper()].value = self._header_value(report, data, cell_map.value)

        # Day header rows, and the days the month does not have.
        rows = [row for row in (self.week_row, self.day_row, self.weekday_row) if row]
        previous_week = None
        for offset in range(MAX_DAYS):
            column = first_day + offset
            letter = get_column_letter(column)
            if offset >= len(days):
                for row in rows:
                    sheet.cell(row=row, column=column).value = None
                sheet.column_dimensions[letter].hidden = True
                continue
            day = days[offset]
            if self.week_row:
                sheet.cell(row=self.week_row, column=column).value = \
                    day['week'] if day['week'] != previous_week else None
                previous_week = day['week']
            if self.day_row:
                sheet.cell(row=self.day_row, column=column).value = day['day']
            if self.weekday_row:
                sheet.cell(row=self.weekday_row, column=column).value = day['weekday_label']
            if self.color_days:
                color = self._day_color(day)
                if color:
                    fill = PatternFill('solid', start_color=color)
                    for row in rows:
                        sheet.cell(row=row, column=column).fill = fill
                    for first, last in filled_rows:
                        for row in range(first, last + 1):
                            sheet.cell(row=row, column=column).fill = fill

        # The blocks were resized: the print area follows the sheet's used
        # range (hidden day columns are not printed).
        sheet.print_area = "A1:%s%s" % (get_column_letter(sheet.max_column), sheet.max_row)

        output = io.BytesIO()
        book.save(output)
        return output.getvalue()

    def _day_color(self, day):
        if day['weekend']:
            return self.weekend_color
        if day['holiday']:
            return self.holiday_color
        if day['absence'] and self.kind == 'internal':
            return self.absence_color
        return False

    @api.model
    def _line_value(self, line, value):
        if value == 'label':
            return line.get('label') or None
        if value == 'client':
            return line.get('partner') or None
        if value == 'project':
            return line.get('project') or None
        if value == 'leave_type':
            return line.get('label') if 'leave_type_id' in line else None
        if value == 'total':
            return line.get('total')
        return None

    def _header_value(self, report, data, value):
        return {
            'employee': data['employee'],
            'company': data['company'],
            'month': fields.Date.to_date(data['date_from']),
            'month_label': data['month_label'],
            'potential_days': data['potential_days'],
            'status': data['status'],
            'submit_date': report.submit_date or None,
            'validate_date': report.validate_date or None,
            'client': data.get('partner') or None,
            'mission_total': data['mission_total'],
            'absence_total': data['absence_total'] if self.kind == 'internal' else None,
            'total': data['mission_total'] + (data['absence_total'] if self.kind == 'internal' else 0),
        }.get(value)

    @api.model
    def _resize_block(self, sheet, first, last, count):
        """Make the block span exactly ``count`` rows, from ``first``.

        Rows are added or removed at the end of the block. openpyxl moves the
        cells but neither row heights, merged cells, nor the ranges of the
        formulas further down: they are shifted here, so the total line keeps
        its look and sums exactly the block's lines. New rows take the style
        of the block's first row.
        """
        delta = count - (last - first + 1)
        if not delta:
            return
        body = {c.column: copy(c._style) for c in sheet[first] if c.has_style}
        height = sheet.row_dimensions[first].height
        new_last = last + delta
        below = {row: (dim.height, dim.hidden)
                 for row, dim in sheet.row_dimensions.items() if row > last}
        if delta > 0:
            sheet.insert_rows(last + 1, delta)
        else:
            sheet.delete_rows(new_last + 1, -delta)
        for row in [row for row in sheet.row_dimensions if row > new_last]:
            del sheet.row_dimensions[row]
        for row, (row_height, hidden) in below.items():
            dimension = sheet.row_dimensions[row + delta]
            dimension.height = row_height
            dimension.hidden = hidden
        for row in range(last + 1, new_last + 1):
            sheet.row_dimensions[row].height = height
            for column, style in body.items():
                sheet.cell(row=row, column=column)._style = copy(style)

        for merged in list(sheet.merged_cells.ranges):
            if merged.min_row > last:
                merged.shift(0, delta)
            elif merged.max_row > new_last:
                sheet.merged_cells.remove(merged)

        def shift(match):
            start = int(match.group(2))
            text = "%s%s" % (match.group(1), start + delta if start > last else start)
            if match.group(3):
                end = int(match.group(4))
                if end > last:
                    end += delta
                elif end == last and start <= last:
                    end = new_last  # range down to the bottom of the block
                text += ":%s%s" % (match.group(3), end)
            return text

        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    cell.value = CELL_REF_RE.sub(shift, cell.value)


class MissionActivityReportTemplateCell(models.Model):
    _name = 'mission.activity.report.template.cell'
    _description = "Activity report template mapping"
    _order = 'kind desc, id'

    template_id = fields.Many2one(
        'mission.activity.report.template', required=True, ondelete='cascade')
    cell = fields.Char(
        string="Column or cell", required=True,
        help="A column letter (“A”): the value of every line of the block. "
             "A cell (“B3”): a header value, written once. For a merged cell, "
             "its top-left cell.")
    kind = fields.Selection(
        [('line', "Lines"), ('header', "Header")], compute='_compute_kind', store=True)
    block = fields.Selection(
        [('mission', "Missions"), ('absence', "Time off"), ('both', "Both")],
        default='both', help="For a column: the block(s) of lines it applies to.")
    value = fields.Selection(VALUES, required=True)

    @api.depends('cell')
    def _compute_kind(self):
        for cell_map in self:
            match = CELL_RE.match((cell_map.cell or '').strip().upper())
            cell_map.kind = 'header' if match and match.group(2) else 'line'

    @api.constrains('cell', 'value')
    def _check_cell(self):
        line_keys = {key for key, _label in LINE_VALUES}
        header_keys = {key for key, _label in HEADER_VALUES}
        for cell_map in self:
            if not CELL_RE.match((cell_map.cell or '').strip().upper()):
                raise UserError(_("“%s”: a column letter (A) or a cell (B3).", cell_map.cell))
            allowed = header_keys if cell_map.kind == 'header' else line_keys
            if cell_map.value not in allowed:
                raise UserError(_(
                    "“%(value)s” cannot go in %(cell)s.",
                    value=dict(VALUES)[cell_map.value], cell=cell_map.cell))


class MissionActivityReportExport(models.TransientModel):
    _name = 'mission.activity.report.export'
    _description = "Export activity reports to Excel"

    report_ids = fields.Many2many('mission.activity.report', string="Reports", required=True)
    template_id = fields.Many2one(
        'mission.activity.report.template', string="Template", required=True,
        default=lambda self: self.env['mission.activity.report.template'].search([], limit=1))

    def action_export(self):
        """One workbook per report (per report and client for a client
        template); several workbooks come as a zip archive."""
        self.ensure_one()
        files = []
        for report in self.report_ids:
            stem = re.sub(r'[\\/:*?"<>|]+', '-', "%s - %s" % (self.template_id.name, report.name))
            if self.template_id.kind == 'client':
                for partner_id in report._get_client_partner_ids():
                    partner = self.env['res.partner'].browse(partner_id).display_name \
                        if partner_id else _("No client")
                    files.append(("%s - %s.xlsx" % (stem, re.sub(r'[\\/:*?"<>|]+', '-', partner)),
                                  self.template_id._export(report, partner_id)))
            else:
                files.append(("%s.xlsx" % stem, self.template_id._export(report)))
        if not files:
            raise UserError(_("These reports hold no mission to export."))
        if len(files) == 1:
            name, content = files[0]
            mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        else:
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
                for file_name, file_content in files:
                    archive.writestr(file_name, file_content)
            name, content, mimetype = _("Activity reports.zip"), buffer.getvalue(), 'application/zip'
        attachment = self.env['ir.attachment'].create({
            'name': name,
            'raw': content,
            'mimetype': mimetype,
            'res_model': self._name,
            'res_id': self.id,
        })
        # target 'new' with close: the download starts and the export dialog
        # closes ('self' left the dialog open over the report).
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'new',
            'close': True,
        }
