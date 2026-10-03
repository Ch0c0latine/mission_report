# -*- coding: utf-8 -*-
import io
from datetime import date, datetime

from freezegun import freeze_time

from odoo import Command
from odoo.exceptions import UserError, ValidationError
from odoo.tests.common import TransactionCase, new_test_user

from odoo.addons.mission_report.models.activity_report import INTERNAL, INTERNAL_KEY
from odoo.addons.mission_report.models.public_holiday_wizard import mission_timezone

# Même principe que les autres tests : 2031 ne porte aucune saisie réelle.
YEAR = 2031


class _Stub:
    """Un calendrier, un salarié ou une société réduits à ce que lit mission_timezone."""

    def __init__(self, **values):
        self.__dict__.update(values)

    def sudo(self):
        return self


class TestRevueCorrections(TransactionCase):
    """Corrections issues de la relecture du code."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env.company
        cls.user = new_test_user(cls.env, 'revue_employee_mission_report', groups='base.group_user')
        cls.manager = new_test_user(
            cls.env, 'revue_manager_mission_report',
            groups='base.group_user,hr_holidays.group_hr_holidays_user')
        cls.employee = cls.env['hr.employee'].create({
            'name': 'Revue Employee', 'company_id': cls.company.id,
            'user_id': cls.user.id, 'leave_manager_id': cls.manager.id})
        cls.client = cls.env['res.partner'].create({'name': 'Revue Client'})
        cls.project = cls.env['project.project'].create({
            'name': 'Revue Project', 'partner_id': cls.client.id})
        cls.env['project.task'].create({
            'name': 'Task', 'project_id': cls.project.id, 'user_ids': [Command.set(cls.user.ids)]})

    # ------------------------------------------------------------------
    # 1. Circuit de validation
    # ------------------------------------------------------------------

    def _employee_report(self, month='2031-08-01'):
        return self.env['mission.activity.report'].with_user(self.user).create({
            'employee_id': self.employee.id, 'date_from': month, 'state': 'draft'})

    def test_employee_creates_a_draft_report(self):
        report = self._employee_report('2031-08-20')
        self.assertEqual(report.state, 'draft')
        self.assertEqual(report.date_from, date(2031, 8, 1))

    def test_employee_cannot_create_a_validated_report(self):
        Report = self.env['mission.activity.report'].with_user(self.user)
        for state in ('validated', 'submitted'):
            with self.assertRaises(UserError):
                Report.create({'employee_id': self.employee.id, 'date_from': '2031-08-05',
                               'state': state, 'snapshot': {'forged': True}})
        self.assertFalse(Report.sudo().search([('employee_id', '=', self.employee.id)]))

    def test_employee_cannot_create_with_a_workflow_field(self):
        Report = self.env['mission.activity.report'].with_user(self.user)
        for vals in ({'snapshot': {'forged': True}},
                     {'submit_date': '2031-08-01'},
                     {'submitted_by_id': self.user.id},
                     {'validate_date': '2031-08-01'},
                     {'validated_by_id': self.manager.id}):
            with self.assertRaises(UserError, msg=str(vals)):
                Report.create({'employee_id': self.employee.id, 'date_from': '2031-08-05', **vals})

    def test_client_cannot_forge_the_workflow_flag(self):
        report = self._employee_report()
        forged = (
            {'mission_report_workflow': True},
            {INTERNAL_KEY: True},
            {INTERNAL_KEY: 'guess'},
            {'mission_report_workflow': True, INTERNAL_KEY: 'guess'},
        )
        for context in forged:
            with self.assertRaises(UserError, msg=str(context)):
                report.with_context(**context).write({'state': 'validated', 'snapshot': {'forged': True}})
            with self.assertRaises(UserError, msg=str(context)):
                self.env['mission.activity.report'].with_user(self.user).with_context(**context).create({
                    'employee_id': self.employee.id, 'date_from': '2031-09-01', 'state': 'validated'})
        self.assertEqual(report.state, 'draft')
        self.assertFalse(report.snapshot)
        with self.assertRaises(UserError):
            report.write({'snapshot': {'forged': True}})

    def test_workflow_buttons_still_write(self):
        report = self._employee_report()
        report.action_submit()
        self.assertEqual(report.state, 'submitted')
        self.assertEqual(report.submitted_by_id, self.user)
        report.with_user(self.manager).action_validate()
        self.assertEqual(report.state, 'validated')
        self.assertEqual(report.validated_by_id, self.manager)
        self.assertTrue(report.snapshot)
        report.with_user(self.manager).action_reset_to_draft()
        self.assertEqual(report.state, 'draft')
        self.assertFalse(report.snapshot)

    def test_internal_token_is_not_the_old_flag(self):
        self.assertNotIn(INTERNAL, (True, 'True', 'true', 1, '1', ''))
        self.assertGreaterEqual(len(INTERNAL), 16)

    # ------------------------------------------------------------------
    # 2. IGD hors périmètre
    # ------------------------------------------------------------------

    @classmethod
    def _igd_people(cls):
        Employee = cls.env['hr.employee']
        approver = new_test_user(
            cls.env, 'revue_approver_mission_report',
            groups='base.group_user,hr_expense.group_hr_expense_team_approver')
        own = Employee.create({'name': 'Approver', 'user_id': approver.id, 'company_id': cls.company.id})
        stranger = Employee.create({'name': 'Stranger', 'company_id': cls.company.id})
        # Le parent est posé : sinon il se déduit du responsable du département.
        subordinate = Employee.create({'name': 'Subordinate', 'parent_id': own.id,
                                       'company_id': cls.company.id})
        department = cls.env['hr.department'].create({'name': 'Revue department', 'manager_id': own.id})
        in_department = Employee.create({'name': 'In department', 'department_id': department.id,
                                         'parent_id': stranger.id, 'company_id': cls.company.id})
        managed = Employee.create({'name': 'Expenses managed', 'parent_id': stranger.id,
                                   'expense_manager_id': approver.id, 'company_id': cls.company.id})
        return approver, own, stranger, subordinate, in_department, managed

    def test_team_approver_is_limited_to_the_rule_domain(self):
        approver, own, stranger, subordinate, in_department, managed = self._igd_people()
        everyone = own | stranger | subordinate | in_department | managed
        wizard = self.env['mission.igd.wizard'].with_user(approver).create({
            'report_month': '5', 'report_year': str(YEAR)})
        allowed = wizard._igd_allowed_employees(everyone.with_user(approver))
        self.assertEqual(set(allowed.ids), set((own | subordinate | in_department | managed).ids))
        self.assertNotIn(stranger.id, allowed.ids)

    def test_team_approver_cannot_generate_for_a_stranger(self):
        approver, own, stranger, *_others = self._igd_people()
        wizard = self.env['mission.igd.wizard'].create({
            'employee_ids': [Command.set(stranger.ids)], 'report_month': '5', 'report_year': str(YEAR)})
        with self.assertRaises(UserError) as caught:
            wizard.with_user(approver).action_generate()
        self.assertIn('Stranger', str(caught.exception))
        self.assertFalse(self.env['hr.expense'].search([('employee_id', '=', stranger.id)]))

    def test_employee_generates_only_their_own(self):
        approver, own, stranger, *_others = self._igd_people()
        wizard = self.env['mission.igd.wizard'].with_user(self.user).create({
            'report_month': '5', 'report_year': str(YEAR)})
        everyone = (own | stranger | self.employee).with_user(self.user)
        self.assertEqual(wizard._igd_allowed_employees(everyone).ids, self.employee.ids)

    def test_expense_officer_keeps_everything(self):
        approver, own, stranger, subordinate, in_department, managed = self._igd_people()
        officer = new_test_user(
            self.env, 'revue_officer_mission_report',
            groups='base.group_user,hr_expense.group_hr_expense_user')
        everyone = own | stranger | subordinate | in_department | managed
        wizard = self.env['mission.igd.wizard'].with_user(officer).create({
            'report_month': '5', 'report_year': str(YEAR)})
        self.assertEqual(set(wizard._igd_allowed_employees(everyone.with_user(officer)).ids),
                         set(everyone.ids))

    # ------------------------------------------------------------------
    # 3. Fuseau horaire
    # ------------------------------------------------------------------

    def test_timezone_chain(self):
        company_calendar = _Stub(tz='Asia/Tokyo')
        company = _Stub(resource_calendar_id=company_calendar)
        employee_calendar = _Stub(tz='Pacific/Auckland')
        # Calendrier de l'employé d'abord, même si l'employé a un autre fuseau.
        employee = _Stub(resource_calendar_id=employee_calendar, tz='America/New_York')
        self.assertEqual(mission_timezone(company, employee).zone, 'Pacific/Auckland')
        # Sans calendrier propre : celui de la société.
        employee = _Stub(resource_calendar_id=False, tz='America/New_York')
        self.assertEqual(mission_timezone(company, employee).zone, 'Asia/Tokyo')
        # Un calendrier sans fuseau : le fuseau de l'employé.
        employee = _Stub(resource_calendar_id=_Stub(tz=False), tz='America/New_York')
        self.assertEqual(mission_timezone(company, employee).zone, 'America/New_York')
        # Sans salarié, le calendrier de la société.
        self.assertEqual(mission_timezone(company).zone, 'Asia/Tokyo')
        # Rien du tout : UTC.
        self.assertEqual(mission_timezone(_Stub(resource_calendar_id=False)).zone, 'UTC')
        self.assertEqual(mission_timezone(_Stub(resource_calendar_id=False),
                                          _Stub(resource_calendar_id=False, tz=False)).zone, 'UTC')

    def test_public_holidays_are_read_in_the_calendar_timezone(self):
        calendar = self.env['resource.calendar'].create({'name': 'Revue Auckland', 'tz': 'Pacific/Auckland'})
        employee = self.env['hr.employee'].create({
            'name': 'Auckland Employee', 'company_id': self.company.id,
            'resource_calendar_id': calendar.id})
        self.assertEqual(mission_timezone(self.company, employee).zone, 'Pacific/Auckland')
        # 14 juillet à Auckland (UTC+12 en hiver) : du 13 à 12 h au 14 à 11 h 59 UTC.
        self.env['resource.calendar.leaves'].create({
            'name': 'Bastille Day', 'company_id': self.company.id, 'calendar_id': False,
            'date_from': datetime(YEAR, 7, 13, 12, 0, 0), 'date_to': datetime(YEAR, 7, 14, 11, 59, 59),
            'time_type': 'leave'})
        holidays = self.env['mission.activity.report']._get_public_holidays(
            employee.sudo(), date(YEAR, 7, 1), date(YEAR, 7, 31))
        self.assertEqual(list(holidays), [date(YEAR, 7, 14)],
                         "lu dans le fuseau de New York, le 13 juillet serait aussi férié")

    def test_holiday_wizard_uses_the_company_calendar_timezone(self):
        self.company.resource_calendar_id.tz = 'Pacific/Auckland'
        self.env.user.tz = 'America/New_York'
        created = self.env['mission.public.holiday.wizard'].create({
            'year': YEAR + 2, 'company_id': self.company.id})._generate()
        bastille = created.filtered(lambda leave: (leave.date_from.month, leave.date_from.day) == (7, 13))
        self.assertEqual(len(bastille), 1)
        self.assertEqual(bastille.date_from, datetime(YEAR + 2, 7, 13, 12, 0, 0))

    # ------------------------------------------------------------------
    # 4. Injection de formule dans l'export Excel
    # ------------------------------------------------------------------

    def test_excel_export_never_writes_a_formula_from_data(self):
        import openpyxl
        names = {
            'employee': '=1+1',
            'client': '=HYPERLINK("http://example.com","x")',
            'project': '@SUM(1+1)',
        }
        employee = self.env['hr.employee'].create({
            'name': names['employee'], 'company_id': self.company.id, 'user_id': self.user.id,
            'leave_manager_id': self.manager.id})
        client = self.env['res.partner'].create({'name': names['client']})
        project = self.env['project.project'].create({'name': names['project'], 'partner_id': client.id})
        self.env['project.task'].create({
            'name': 'Task', 'project_id': project.id, 'user_ids': [Command.set(self.user.ids)]})
        self.env['hr.leave'].create({
            'employee_id': employee.id, 'project_id': project.id,
            'request_date_from': '2031-10-06', 'request_date_to': '2031-10-07'})
        report = self.env['mission.activity.report'].create({
            'employee_id': employee.id, 'date_from': '2031-10-01'})
        template = self.env.ref('mission_report.activity_report_template_internal')
        template.action_generate_file()
        sheet = openpyxl.load_workbook(io.BytesIO(template._export(report))).worksheets[0]
        types = {}
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    types.setdefault(cell.value, set()).add(cell.data_type)
        for key, text in names.items():
            self.assertIn(text, types, "%s est écrit dans le fichier" % key)
            self.assertEqual(types[text], {'s'}, "%s reste un texte" % key)
        # Les totaux du modèle, eux, restent des formules.
        self.assertTrue(any(value.startswith('=SUM(') for value in types))

    def test_set_text_covers_every_prefix(self):
        import openpyxl
        from odoo.addons.mission_report.models.activity_report_template import _set_text
        sheet = openpyxl.Workbook().active
        for index, text in enumerate(('=1+1', '+1', '-1+1', '@A1', 'Client'), start=1):
            cell = sheet.cell(row=index, column=1)
            _set_text(cell, text)
            self.assertEqual(cell.data_type, 's', text)
            self.assertEqual(cell.value, text)
        cell = sheet.cell(row=6, column=1)
        _set_text(cell, 1.5)
        self.assertEqual(cell.data_type, 'n')
        _set_text(cell, None)
        self.assertIsNone(cell.value)

    # ------------------------------------------------------------------
    # 5. Recherches de tâches en droits utilisateur
    # ------------------------------------------------------------------

    def test_manager_sees_the_missions_of_another_employee(self):
        # Projet à l'invitation : la règle de visibilité des tâches les cache au responsable.
        project = self.env['project.project'].create({
            'name': 'Revue Private', 'partner_id': self.client.id, 'privacy_visibility': 'followers'})
        self.env['project.task'].create({
            'name': 'Private task', 'project_id': project.id, 'user_ids': [Command.set(self.user.ids)]})
        Leave = self.env['hr.leave'].with_user(self.manager)
        employee = self.employee.with_user(self.manager)
        projects = Leave._mission_projects_for_period(employee, date(YEAR, 8, 1), date(YEAR, 8, 31))
        self.assertIn(project.id, projects.ids)
        available = Leave.new({'employee_id': employee.id}).available_project_ids
        self.assertIn(project.id, available.ids)

    def test_manager_can_enter_a_mission_for_another_employee(self):
        project = self.env['project.project'].create({
            'name': 'Revue Private 2', 'partner_id': self.client.id, 'privacy_visibility': 'followers'})
        self.env['project.task'].create({
            'name': 'Private task', 'project_id': project.id, 'user_ids': [Command.set(self.user.ids)]})
        leave = self.env['hr.leave'].create({
            'employee_id': self.employee.id, 'project_id': project.id,
            'request_date_from': '2031-08-04', 'request_date_to': '2031-08-04'})
        # Ne lève pas « n'est assigné(e) à aucune tâche » à tort.
        leave.with_user(self.manager)._check_employee_assigned_to_project()

    def test_unassigned_employee_is_still_refused(self):
        other_user = new_test_user(self.env, 'revue_other_mission_report', groups='base.group_user')
        other = self.env['hr.employee'].create({
            'name': 'Not assigned', 'company_id': self.company.id, 'user_id': other_user.id})
        with self.assertRaises(ValidationError):
            self.env['hr.leave'].create({
                'employee_id': other.id, 'project_id': self.project.id,
                'request_date_from': '2031-08-04', 'request_date_to': '2031-08-04'})

    # ------------------------------------------------------------------
    # 6. Valeurs par défaut : la date du fuseau de l'utilisateur
    # ------------------------------------------------------------------

    @freeze_time('2031-06-30 23:30:00')
    def test_defaults_follow_the_user_timezone(self):
        # 1er juillet 1 h 30 à Paris, encore le 30 juin en UTC.
        Report = self.env['mission.activity.report'].with_context(tz='Europe/Paris')
        self.assertEqual(Report.default_get(['date_from'])['date_from'], date(YEAR, 7, 1))
        Igd = self.env['mission.igd.wizard'].with_context(tz='Europe/Paris')
        defaults = Igd.default_get(['report_month', 'report_year'])
        self.assertEqual((defaults['report_month'], defaults['report_year']), ('7', str(YEAR)))
        # Le même instant, vu d'un fuseau en retard sur UTC.
        Report = self.env['mission.activity.report'].with_context(tz='America/New_York')
        self.assertEqual(Report.default_get(['date_from'])['date_from'], date(YEAR, 6, 1))

    @freeze_time('2031-12-31 23:30:00')
    def test_holiday_wizard_year_follows_the_user_timezone(self):
        Wizard = self.env['mission.public.holiday.wizard']
        self.assertEqual(Wizard.with_context(tz='Europe/Paris').default_get(['year'])['year'], YEAR + 1)
        self.assertEqual(Wizard.with_context(tz='America/New_York').default_get(['year'])['year'], YEAR)


class TestMissionPeriodCheck(TransactionCase):
    """Le contrôle des périodes d'affaires ne se coupe pas depuis le client."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        client = cls.env['res.partner'].create({'name': 'Period Client'})
        cls.project = cls.env['project.project'].create({'name': 'Period Project', 'partner_id': client.id})
        service = cls.env['product.product'].create({
            'name': 'Period days', 'type': 'service', 'list_price': 500.0, 'invoice_policy': 'delivery'})

        def confirmed_order():
            order = cls.env['sale.order'].create({
                'partner_id': client.id,
                'order_line': [Command.create({'product_id': service.id, 'product_uom_qty': 10.0})]})
            order.action_confirm()
            return order

        cls.first = confirmed_order()
        cls.first.write({'project_id': cls.project.id,
                         'mission_date_start': '2031-01-01', 'mission_date_end': '2031-05-12'})
        cls.second = confirmed_order()
        cls.second.write({'project_id': cls.project.id,
                          'mission_date_start': '2031-05-13', 'mission_date_end': '2031-12-31'})

    def test_overlap_is_refused_whatever_the_context(self):
        for context in ({}, {'mission_no_check': True}, {INTERNAL_KEY: True}, {INTERNAL_KEY: 'guess'}):
            with self.assertRaises(ValidationError, msg=str(context)):
                self.second.with_context(**context).write({'mission_date_start': '2031-05-01'})
        self.assertEqual(self.second.mission_date_start, date(YEAR, 5, 13))

    def test_internal_write_skips_the_check(self):
        self.second.with_context(**{INTERNAL_KEY: INTERNAL}).write({'mission_date_start': '2031-05-01'})
        self.assertEqual(self.second.mission_date_start, date(YEAR, 5, 1))

    def test_dates_are_filled_from_the_volet_title_by_the_module(self):
        third = self.env['sale.order'].create({
            'partner_id': self.project.partner_id.id, 'project_id': self.project.id,
            'note': "Volet 3 : période du 01/01/2032 au 31/12/2032"})
        self.assertEqual(third.mission_date_start, date(2032, 1, 1))
        self.assertEqual(third.mission_date_end, date(2032, 12, 31))
